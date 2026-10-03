from __future__ import annotations
from datetime import datetime
from pathlib import Path
from openpyxl import load_workbook
from ..models import Activity, Relationship, Schedule

def _dt(v):
    if isinstance(v, datetime): return v
    if not v: return None
    try: return datetime.fromisoformat(str(v))
    except ValueError: return None

def parse_excel(path: str|Path, column_map: dict[str,str], activity_sheet="Activities", relationship_sheet="Relationships") -> Schedule:
    wb=load_workbook(path, data_only=True, read_only=True); ws=wb[activity_sheet]
    rows=ws.iter_rows(values_only=True); headers=[str(x).strip() if x is not None else "" for x in next(rows)]
    idx={h:i for i,h in enumerate(headers)}
    def val(row,key,default=None):
        col=column_map.get(key)
        return row[idx[col]] if col in idx else default
    activities={}
    for row in rows:
        aid=str(val(row,"id","")).strip()
        if not aid: continue
        activities[aid]=Activity(id=aid,name=str(val(row,"name","")).strip(),duration=float(val(row,"duration",0) or 0),percent_complete=float(val(row,"percent_complete",0) or 0),actual_start=_dt(val(row,"actual_start")),actual_finish=_dt(val(row,"actual_finish")),planned_start=_dt(val(row,"planned_start")),planned_finish=_dt(val(row,"planned_finish")),wbs=str(val(row,"wbs","") or "") or None)
    rels=[]
    if relationship_sheet in wb.sheetnames:
        rw=wb[relationship_sheet]; rr=rw.iter_rows(values_only=True); rh=[str(x).strip() if x is not None else "" for x in next(rr)]; ri={h:i for i,h in enumerate(rh)}
        for row in rr:
            p=str(row[ri[column_map["predecessor_id"]]]).strip(); s=str(row[ri[column_map["successor_id"]]]).strip()
            typ=str(row[ri[column_map.get("relationship_type","Type")]] or "FS").strip()
            lag=float(row[ri[column_map.get("lag","Lag")]] or 0)
            rels.append(Relationship(p,s,typ,lag))
    return Schedule(Path(path).stem,Path(path).stem,activities,rels,source_type="Excel")
